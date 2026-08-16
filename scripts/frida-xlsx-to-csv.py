#!/usr/bin/env python3
"""Extract a flat CSV from DTU Frida's published .xlsx, for convert-frida.ts.

Frida's `Data_Table` sheet can't be exported to usable CSV by hand: it has FOUR
header rows (Danish parameter names, English parameter names, units, parameter
ids) before the data starts, and the food ids live on a separate `Food` sheet
that has to be joined on the Danish food name. This script does exactly that
join and nothing else.

    python scripts/frida-xlsx-to-csv.py Frida_5.5_Dataset.xlsx frida.csv
    npm run frida:convert -- frida.csv ./data/frida.json --version 5.5

Requires: openpyxl (pip install openpyxl). This is the one manual prerequisite
in an otherwise Node-only repo — Frida ships .xlsx/.ods only, with no API.
"""
import csv
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# Column indices in Data_Table, identified from the English header row (row 2).
# Verified against Frida 5.5; re-check if DTU reorders columns in a new version.
COLUMNS = [
    (5, "Energy (kcal)"),
    (7, "Protein"),
    (11, "Available carbohydrates"),  # not "by difference" (col 10) — this is the labelling figure
    (14, "Fat"),
    (13, "Dietary fibre"),
    (97, "Sum sugars"),
    (180, "Sum saturated fatty acids"),
    (16, "Salt labelling"),
]
DATA_STARTS_AT_ROW = 5
ENGLISH_HEADER_ROW = 2


def main() -> None:
    if len(sys.argv) < 3:
        sys.exit("Usage: frida-xlsx-to-csv.py <Frida_x.x_Dataset.xlsx> <output.csv>")
    src, dst = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    # Sanity-check our hard-coded indices against the sheet's own English header
    # row, so a reordered future release fails loudly instead of silently
    # writing the wrong nutrient into the wrong column.
    header = list(wb["Data_Table"].iter_rows(min_row=ENGLISH_HEADER_ROW, max_row=ENGLISH_HEADER_ROW, values_only=True))[0]
    for index, expected in COLUMNS:
        actual = (header[index] or "").strip()
        if actual != expected:
            sys.exit(f"Column {index} is '{actual}', expected '{expected}'. Frida's layout changed — update COLUMNS.")

    # FoodID lives on a separate sheet; join on the Danish name.
    ids = {}
    for row in wb["Food"].iter_rows(min_row=2, values_only=True):
        danish_name, _english, food_id = row[0], row[1], row[2]
        if danish_name and food_id:
            ids[str(danish_name).strip()] = str(food_id).strip()

    written = skipped = 0
    with open(dst, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["FoodID", "Food name"] + [name for _, name in COLUMNS])
        for row in wb["Data_Table"].iter_rows(min_row=DATA_STARTS_AT_ROW, values_only=True):
            name = (row[0] or "").strip() if row[0] else ""
            if not name:
                continue
            food_id = ids.get(name)
            if not food_id:
                skipped += 1
                continue
            writer.writerow([food_id, name] + [row[i] if row[i] is not None else "" for i, _ in COLUMNS])
            written += 1

    print(f"Wrote {written} foods to {dst}" + (f" ({skipped} skipped: no FoodID match)" if skipped else ""))


if __name__ == "__main__":
    main()
